from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = r"C:\Users\ASUS\Documents\Unity\Audio_procedural20261\Proyecto_Sintesis_Procedural_Groq_Wavetables.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_style(table, header_fill="1F4E79"):
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6FA")


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.add_run(f"{index}.  {item}")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08


def add_code(doc, text):
    for line in text.strip("\n").splitlines():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8.7)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.italic = True
    run.font.size = Pt(9)


def set_document_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Title"
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def add_toc(doc):
    add_heading(doc, "Tabla de contenido", 1)
    entries = [
        "1 Objetivo y alcance",
        "2 Arquitectura general del proyecto",
        "3 Funcionamiento de las wavetables multiframe",
        "4 ADSR y transiciones tímbricas",
        "5 Antialiasing por límite de Nyquist",
        "6 Integración con GroqCloud",
        "7 Obtención de la API key de Groq paso a paso",
        "8 Configuración de la API key en Unity",
        "9 Pruebas recomendadas",
        "10 Referencias oficiales",
    ]
    for entry in entries:
        add_body(doc, entry)
    doc.add_page_break()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    set_document_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Proyecto de Síntesis Procedural con Groq y Wavetables Multiframe")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Documento técnico y guía de configuración para estudiantes")
    run.font.size = Pt(13)
    run.font.bold = True
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Desarrollado por: Miguel Ángel Olivares Tenorio").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    repo_run = p.add_run("Repositorio del proyecto: https://github.com/ingmiguelolivares/Audio_procedural20261")
    repo_run.font.size = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Versión del documento: Octubre de 2026")
    doc.add_page_break()

    add_body(
        doc,
        "Este documento explica el funcionamiento del sintetizador procedural, la integración con GroqCloud "
        "y el modelo de wavetables multiframe implementado para producir timbres generados por inteligencia artificial. "
        "Está dirigido a estudiantes que ya tienen el proyecto abierto en Unity y necesitan comprender cómo viaja un "
        "timbre desde una descripción textual hasta las voces audibles del sintetizador."
    )
    add_body(
        doc,
        "La idea central es mantener el proyecto como un sistema de audio procedural. GroqCloud interpreta una descripción "
        "sonora y devuelve parámetros estructurados; Unity construye localmente la wavetable, aplica el ADSR y reproduce el "
        "audio mediante los motores existentes. La IA no genera archivos de audio ni muestras PCM completas."
    )

    add_toc(doc)

    add_heading(doc, "1 Objetivo y alcance", 1)
    add_body(
        doc,
        "El proyecto permite estudiar síntesis procedural con dos rutas de reproducción conservadas y mejoradas: "
        "Legacy moderno y External DSP WebGL. Ambas usan la misma descripción conceptual del timbre, pero pueden reproducirla "
        "con mecanismos internos distintos según la plataforma."
    )
    add_bullets(
        doc,
        [
            "Comprender cómo se representa un timbre mediante armónicos y ADSR.",
            "Usar GroqCloud para convertir prompts en parámetros de síntesis, no en audio terminado.",
            "Reproducir el mismo preset en Legacy moderno y en la ruta External DSP WebGL.",
            "Configurar una API key personal de Groq sin guardarla dentro del código fuente.",
            "Probar el sistema con un preset local aunque no haya Internet o no exista API key.",
        ],
    )

    add_heading(doc, "2 Arquitectura general del proyecto", 1)
    add_body(
        doc,
        "La arquitectura separa tres responsabilidades: interpretación de intención sonora, construcción local del timbre "
        "y reproducción de audio. Esta separación evita crear un sintetizador paralelo y mantiene el flujo pedagógico claro."
    )
    add_code(
        doc,
        """
Prompt del estudiante
        |
        v
GroqTimbreGenerator
        |
        v
JSON estructurado de GroqCloud
        |
        v
TimbrePreset + Wavetable + ADSR
        |
        v
Polifonia
        |
        +--> Voces Osc activas
        |
        +--> Voces Osc futuras creadas por NoteOn
        |
        v
Legacy moderno o External DSP WebGL
        |
        v
Audio audible
""",
    )
    add_caption(doc, "Figura 1 Flujo general de generación y reproducción del timbre")

    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Componente"
    table.rows[0].cells[1].text = "Responsabilidad"
    table.rows[0].cells[2].text = "Archivo principal"
    rows = [
        ("GroqTimbreGenerator", "Solicita JSON estructurado a Groq y construye un resultado de timbre.", "Assets/scripts/TimbreAI/GroqTimbreGenerator.cs"),
        ("TimbrePreset", "Modelo compartido de armónicos multiframe y ADSR.", "Assets/scripts/TimbreAI/TimbrePreset.cs"),
        ("Polifonia", "Propaga el timbre generado a voces activas y futuras.", "Assets/scripts/ProceduralAudio/NativeDSP/Polifonia.cs"),
        ("Osc", "Reproduce una voz y aplica wavetable, ADSR y backend seleccionado.", "Assets/scripts/ProceduralAudio/NativeDSP/Osc.cs"),
        ("External DSP WebGL", "Reproduce las voces fuera de OnAudioFilterRead cuando corresponde.", "ProceduralSynthVoice y ProceduralSynth.jslib"),
    ]
    for component, responsibility, file_name in rows:
        cells = table.add_row().cells
        cells[0].text = component
        cells[1].text = responsibility
        cells[2].text = file_name
    set_table_style(table)

    add_heading(doc, "Aclaración sobre el nombre Gemini", 2)
    add_body(
        doc,
        "En algunas partes internas del proyecto puede aparecer el valor heredado WaveFormType.Gemini. En la implementación "
        "actual no significa que Gemini sea el proveedor de IA. Ese nombre funciona como etiqueta legacy para el modo de "
        "wavetable generada por IA. El proveedor actual documentado aquí es GroqCloud. GroqCloud tampoco es una forma de "
        "onda; es el servicio que ayuda a generar parámetros estructurados."
    )

    add_heading(doc, "3 Funcionamiento de las wavetables multiframe", 1)
    add_body(
        doc,
        "Una wavetable tradicional contiene un ciclo periódico que el oscilador recorre a diferentes velocidades para producir "
        "distintas notas. La implementación actual amplía esta idea con tres frames tímbricos, cada uno representado por hasta "
        "16 amplitudes armónicas."
    )
    table = doc.add_table(rows=1, cols=4)
    headers = ["Frame", "Momento musical", "Contenido", "Uso durante la nota"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        ("Attack", "Inicio", "Armónicos del golpe inicial o entrada del sonido.", "Se interpola hacia Sustain mientras comienza la nota."),
        ("Sustain", "Cuerpo estable", "Armónicos principales del timbre sostenido.", "Es el color central mientras la nota permanece activa."),
        ("Release", "Liberación", "Armónicos del final del sonido.", "Se interpola desde Sustain al soltar la tecla o botón."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_style(table, "2F5597")

    add_body(
        doc,
        "Cada frame no guarda miles de muestras. Guarda una descripción compacta del espectro. A partir de esas amplitudes, "
        "Unity puede reconstruir muestras con síntesis aditiva o construir una wavetable compatible con el motor."
    )
    add_body(doc, "La ecuación conceptual de síntesis aditiva es:")
    add_code(doc, "x[n] = suma de Ak * sin(2 * pi * k * n / N)")
    add_body(
        doc,
        "En esa expresión, n es la posición de muestra, N es el tamaño de la tabla, k es el índice del armónico y Ak es la "
        "amplitud del armónico. Si el primer armónico tiene mucha amplitud, el sonido tendrá una fundamental fuerte. Si los "
        "armónicos altos tienen más presencia, el sonido se percibirá más brillante."
    )

    add_heading(doc, "4 ADSR y transiciones tímbricas", 1)
    add_body(
        doc,
        "El ADSR controla la amplitud del sonido. Las transiciones entre Attack, Sustain y Release controlan el color tímbrico. "
        "Ambas cosas ocurren al mismo tiempo, pero no son lo mismo."
    )
    table = doc.add_table(rows=1, cols=3)
    for i, header in enumerate(["Parámetro", "Unidad", "Significado"]):
        table.rows[0].cells[i].text = header
    rows = [
        ("Attack", "Segundos o milisegundos según el campo", "Tiempo para subir desde silencio hasta el nivel máximo."),
        ("Decay", "Segundos o milisegundos según el campo", "Tiempo para bajar desde el máximo hasta el nivel de sustain."),
        ("Sustain Level", "0 a 1", "Nivel relativo que se mantiene mientras la nota sigue activa."),
        ("Release", "Segundos o milisegundos según el campo", "Tiempo para caer hacia silencio después de NoteOff."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_style(table, "404040")
    add_body(
        doc,
        "Cuando se genera un timbre por IA, el ADSR se propaga junto con la wavetable. Polifonia lo conserva para que las "
        "voces que ya suenan y las voces que se creen después reciban los mismos valores."
    )

    add_heading(doc, "5 Antialiasing por límite de Nyquist", 1)
    add_body(
        doc,
        "El antialiasing implementado se basa en una regla simple: no se debe reproducir un armónico cuya frecuencia supere "
        "la mitad de la frecuencia de muestreo. Esta mitad se conoce como frecuencia de Nyquist."
    )
    add_code(
        doc,
        """
frecuencia maxima = sampleRate / 2
frecuencia del armonico = frecuencia fundamental * indice del armonico

si frecuencia del armonico <= frecuencia maxima:
    usar el armonico
si frecuencia del armonico > frecuencia maxima:
    omitir el armonico
""",
    )
    add_body(
        doc,
        "Esta regla se aplica por nota. En notas graves caben más armónicos dentro del límite. En notas agudas se reducen los "
        "armónicos disponibles para evitar componentes que el sistema digital no puede representar correctamente."
    )

    add_heading(doc, "6 Integración con GroqCloud", 1)
    add_body(
        doc,
        "GroqCloud se usa como intérprete de intención sonora. El estudiante escribe una descripción, por ejemplo: "
        "Warm dark analog bass with fast attack and slow release. El sistema solicita a Groq una respuesta JSON con nombre "
        "del timbre, valores perceptuales, tres frames armónicos y ADSR."
    )
    add_code(
        doc,
        """
{
  "timbreName": "Warm Analog Bass",
  "timbre": {
    "brightness": 0.25,
    "harmonicity": 0.9,
    "noisiness": 0.03,
    "warmth": 0.85,
    "fundamentalStrength": 1.0
  },
  "attackHarmonics": [1.0, 0.42, 0.20, ...],
  "sustainHarmonics": [1.0, 0.28, 0.12, ...],
  "releaseHarmonics": [0.85, 0.18, 0.05, ...],
  "adsr": {
    "attackMs": 10,
    "decayMs": 180,
    "sustainMs": 650,
    "sustainLevel": 0.68,
    "releaseMs": 850
  }
}
""",
    )
    add_body(
        doc,
        "La respuesta no contiene una wavetable de 2048 muestras. Eso sería costoso, difícil de validar y poco pedagógico. "
        "En cambio, contiene parámetros compactos que Unity valida y transforma localmente en sonido."
    )

    add_heading(doc, "7 Obtención de la API key de Groq paso a paso", 1)
    add_body(
        doc,
        "Cada estudiante debe crear su propia API key. La clave identifica solicitudes realizadas a GroqCloud y debe tratarse "
        "como una credencial personal. No debe pegarse en repositorios, capturas públicas ni archivos compartidos."
    )
    add_numbered(
        doc,
        [
            "Abrir el sitio oficial de GroqCloud Console en https://console.groq.com.",
            "Iniciar sesión con una cuenta personal o institucional. Si la cuenta institucional bloquea el acceso, usar una cuenta permitida por el docente o revisar las políticas de la organización.",
            "Entrar a la sección API Keys. También se puede abrir directamente https://console.groq.com/keys.",
            "Si la consola solicita seleccionar o crear un proyecto, crear un proyecto de laboratorio con un nombre claro, por ejemplo sintesis-procedural-lab.",
            "Verificar que el proyecto correcto esté seleccionado en el selector de organización y proyecto de la consola.",
            "Seleccionar Create API Key.",
            "Escribir un nombre descriptivo para la clave, por ejemplo Unity Procedural Synth Lab.",
            "Confirmar la creación y copiar la clave cuando Groq la muestre. Normalmente solo se puede ver completa en ese momento.",
            "Guardar la clave temporalmente en un lugar privado mientras se configura Unity. No publicarla ni enviarla por chat abierto.",
        ],
    )
    add_body(
        doc,
        "Según la documentación oficial de Groq, las API keys pertenecen al proyecto seleccionado en la consola. La documentación "
        "de proyectos también indica que propietarios y usuarios con rol de desarrollador pueden crear o administrar claves."
    )

    add_heading(doc, "8 Configuración de la API key en Unity", 1)
    add_body(
        doc,
        "La implementación actual evita guardar la API key dentro del código o dentro de la escena. En el Editor de Unity, "
        "el Inspector personalizado de GroqTimbreGenerator guarda la clave en EditorPrefs. Ese almacenamiento es local al equipo "
        "del estudiante y no se incluye en Assets, prefabs ni builds."
    )
    add_numbered(
        doc,
        [
            "Abrir la escena de prueba del sintetizador.",
            "Seleccionar el objeto que contiene el componente GroqTimbreGenerator.",
            "En el Inspector, ubicar la sección Groq API Key.",
            "Pegar la API key personal en el campo Nueva clave.",
            "Presionar Actualizar clave local.",
            "Verificar que el Inspector indique que la clave está configurada.",
            "Para borrar la clave del equipo, presionar Eliminar clave local.",
        ],
    )
    add_body(
        doc,
        "En ejecución o en WebGL puede usarse un campo de interfaz conectado a runtimeApiKeyInput. Esa clave vive solo durante "
        "la sesión. En navegador no debe considerarse secreta, porque cualquier clave usada del lado del cliente puede ser "
        "inspeccionada. Para una publicación real se recomienda un intermediario seguro en servidor."
    )

    doc.add_page_break()
    add_heading(doc, "9 Pruebas recomendadas", 1)
    table = doc.add_table(rows=1, cols=4)
    for i, header in enumerate(["Prueba", "Qué verifica", "Resultado esperado", "Observación"]):
        table.rows[0].cells[i].text = header
    rows = [
        ("Preset local", "Funcionamiento sin Internet ni API key.", "Se genera un timbre audible.", "Usar Generate Local Test Timbre."),
        ("Prompt con Groq", "Comunicación con Groq y JSON estructurado.", "El estado indica generación correcta.", "No hacer muchas llamadas repetidas."),
        ("Legacy moderno", "Ruta OnAudioFilterRead moderna.", "La nota suena con el timbre generado.", "Probar notas graves y agudas."),
        ("External DSP WebGL", "Ruta externa y navegador.", "La nota suena tras desbloquear audio.", "Requiere interacción del usuario en WebGL."),
        ("Cambio a Sine", "Regresión de formas normales.", "Sine sigue sonando.", "Luego volver al modo generado."),
        ("Modo generado", "La wavetable no se reemplaza por ceros.", "El timbre generado vuelve a sonar.", "El nombre interno puede aparecer como Gemini."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_style(table, "1F4E79")

    add_heading(doc, "10 Referencias oficiales", 1)
    add_bullets(
        doc,
        [
            "Repositorio del proyecto: https://github.com/ingmiguelolivares/Audio_procedural20261",
            "GroqCloud Console API Keys: https://console.groq.com/keys",
            "Groq Projects documentation: https://console.groq.com/docs/projects",
            "Groq Quickstart: https://console.groq.com/docs/quickstart",
            "Groq API Reference: https://console.groq.com/docs/api-reference",
            "Groq Security Onboarding: https://console.groq.com/docs/production-readiness/security-onboarding",
            "Unity Manual sobre audio y ejecución WebGL: https://docs.unity3d.com/Manual/",
        ],
    )

    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Síntesis de Audio Procedural con Inteligencia Artificial")
    run.font.size = Pt(8)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
