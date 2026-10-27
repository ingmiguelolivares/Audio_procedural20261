from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"C:\Users\ASUS\Documents\Unity\Audio_procedural20261\Implementacion_WebGL_Sintesis_Procedural.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_style(table, header_fill="1F4E79"):
    table.style = "Table Grid"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def clear_style_borders(style):
    p_pr = style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in [("Title", 23), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Title"
        clear_style_borders(style)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        p.add_run(f"{index}.  {item}")


def code(doc, text):
    lines = text.strip("\n").splitlines()
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.right_indent = Cm(0.35)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_together = True
        if index < len(lines) - 1:
            p.paragraph_format.keep_with_next = True
        p_pr = p._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        p_pr.append(shading)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(8.5)


def code_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9.5)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.italic = True
    r.font.size = Pt(9)


def add_simple_table(doc, headers, rows, header_fill="1F4E79"):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_style(table, header_fill)
    return table


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Implementación WebGL de la Síntesis Procedural")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Comparación entre Legacy OnAudioFilterRead y reproducción compatible con navegador")
    r.font.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Desarrollado por: Miguel Ángel Olivares Tenorio").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Repositorio: ")
    add_hyperlink(p, "github.com/ingmiguelolivares/Audio_procedural20261", "https://github.com/ingmiguelolivares/Audio_procedural20261")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Versión del documento: Octubre de 2026")
    doc.add_page_break()

    body(
        doc,
        "Este documento explica cómo se adaptó el sintetizador procedural para funcionar en WebGL sin perder la lógica "
        "matemática que ya existía en la implementación Legacy. El punto central para los estudiantes es que la fórmula "
        "de generación de señal se conserva, pero cambia el entorno que ejecuta el procesamiento de audio."
    )
    body(
        doc,
        "En Legacy, Unity llama a OnAudioFilterRead para que el script escriba muestras directamente en un buffer de audio. "
        "En WebGL, Unity no soporta OnAudioFilterRead en la plataforma Web, por lo que el proyecto traslada la reproducción "
        "procedural hacia una capa JavaScript basada en Web Audio. La misma idea de seno, fase, ADSR y armónicos se mantiene, "
        "pero se ejecuta en el navegador."
    )

    heading(doc, "Tabla de contenido", 1)
    for entry in [
        "1 Problema que resuelve la implementación WebGL",
        "2 Funcionamiento Legacy basado en OnAudioFilterRead",
        "3 Limitación técnica de WebGL",
        "4 Nueva arquitectura compatible con WebGL",
        "5 Cómo se conserva la generación de señales seno",
        "6 Comparación del procesamiento de buffers",
        "7 Puente entre C Sharp y JavaScript",
        "8 ADSR equivalente en ambos motores",
        "9 Wavetables y lectura por fase",
        "10 Timbre multiframe y antialiasing",
        "11 Ciclo de vida de una voz",
        "12 Comparación integral Legacy y WebGL",
        "13 Recorrido completo de una nota",
        "14 Prácticas y pruebas para estudiantes",
        "15 Repositorio y referencias",
    ]:
        body(doc, entry)
    doc.add_page_break()

    heading(doc, "1 Problema que resuelve la implementación WebGL", 1)
    body(
        doc,
        "El framework Legacy del proyecto se diseñó alrededor de OnAudioFilterRead. Ese enfoque funciona bien en el Editor "
        "y en plataformas de escritorio porque Unity entrega un arreglo de muestras que el script puede llenar en tiempo real. "
        "El navegador, sin embargo, usa un modelo diferente de audio y no permite depender de ese mismo callback de Unity."
    )
    bullets(
        doc,
        [
            "Se conserva la generación procedural del sonido.",
            "Se evita depender de archivos de audio pregrabados.",
            "Se mantiene la misma lógica de seno, armónicos, ADSR, wavetable y notas.",
            "Se adapta la salida de audio para que el navegador pueda reproducirla.",
            "Se separa la configuración del timbre en C# de la reproducción final en Web Audio.",
        ],
    )

    heading(doc, "2 Funcionamiento Legacy basado en OnAudioFilterRead", 1)
    body(
        doc,
        "En Legacy, cada voz Osc calcula muestras dentro de OnAudioFilterRead. Unity llama a este método cuando necesita "
        "llenar el buffer de audio. El método recorre el arreglo data, calcula una muestra por frame y copia el resultado "
        "en uno o varios canales."
    )
    code(
        doc,
        """
private void OnAudioFilterRead(float[] data, int channels)
{
    for (int i = 0; i < data.Length; i += channels)
    {
        int currentTime = TimeIndex;
        float envelope = getADSR(currentTime);
        float sample = GenerateWaveSample(FormType, frecuencia, currentTime);
        data[i] = sample * envelope * Level;
        TimeIndex++;
    }
}
""",
    )
    body(
        doc,
        "El ejemplo anterior resume el concepto. El código real del proyecto también contempla detune, vibrato, tremolo, "
        "sampling, wavetables, timbre multiframe y finalización de la envolvente."
    )
    body(doc, "La generación básica de una onda seno se basa en la ecuación:")
    code(doc, "sin(2 * pi * frecuencia * tiempo / sampleRate)")
    body(
        doc,
        "En C# esta idea aparece como una función de seno que recibe frecuencia y tiempo de muestra. La frecuencia determina "
        "qué tan rápido avanza la fase. El tiempo de muestra aumenta en cada paso del buffer."
    )

    doc.add_page_break()
    heading(doc, "3 Limitación técnica de WebGL", 1)
    body(
        doc,
        "La documentación oficial de Unity indica que OnAudioFilterRead no está soportado en la plataforma Web. Unity también "
        "explica que, por limitaciones del entorno Web, la plataforma usa una implementación basada en Web Audio API en lugar "
        "de depender del mismo backend de audio disponible en escritorio."
    )
    add_simple_table(
        doc,
        ["Aspecto", "Editor o escritorio", "WebGL"],
        [
            ("Callback de audio", "OnAudioFilterRead puede llenar buffers.", "OnAudioFilterRead no está soportado en Web."),
            ("Backend", "Unity puede usar su ruta normal de audio.", "Unity se apoya en Web Audio API del navegador."),
            ("Hilos", "Existe un hilo de audio con restricciones.", "El navegador impone reglas propias de ejecución."),
            ("Autoplay", "Normalmente se reproduce al iniciar si la escena lo permite.", "El navegador exige interacción del usuario para desbloquear audio."),
        ],
        "404040",
    )
    body(
        doc,
        "Por esta razón, la solución no consiste en insistir en OnAudioFilterRead dentro del build WebGL. La solución consiste "
        "en preservar la matemática del sintetizador y mover la reproducción a un puente compatible con navegador."
    )

    heading(doc, "4 Nueva arquitectura compatible con WebGL", 1)
    body(
        doc,
        "La implementación mantiene Osc y Polifonia como controladores principales en C#. Cuando se usa el backend externo, "
        "Osc no intenta producir todas las muestras directamente dentro de OnAudioFilterRead para WebGL. En su lugar, sincroniza "
        "una configuración de voz hacia ProceduralSynthVoice."
    )
    code(
        doc,
        """
Osc
  -> construye ProceduralSynthVoiceConfig
  -> envía waveform, ADSR, wavetable, armónicos y flags
  -> ProceduralSynthVoice
  -> WebGL jslib
  -> Web Audio API
  -> salida del navegador
""",
    )
    caption(doc, "Figura 1 Flujo de configuración desde C# hacia WebGL")
    add_simple_table(
        doc,
        ["Archivo", "Papel en la implementación WebGL"],
        [
            ("Assets/scripts/ProceduralAudio/NativeDSP/Osc.cs", "Decide el backend, conserva parámetros de la voz y sincroniza la configuración."),
            ("Assets/scripts/ProceduralAudio/NativeDSP/ProceduralSynthVoiceConfig.cs", "Agrupa los datos que deben viajar al backend externo."),
            ("Assets/scripts/ProceduralAudio/NativeDSP/ProceduralSynthVoice.cs", "Puente C# que llama funciones nativas o WebGL según plataforma."),
            ("Assets/scripts/ProceduralAudio/NativeDSP/ManagedProceduralSynthVoice.cs", "Fallback administrado y referencia de la lógica DSP en C#."),
            ("Assets/Plugins/WebGL/ProceduralSynth.jslib", "Implementación JavaScript que genera audio en el navegador."),
        ],
        "1F4E79",
    )

    heading(doc, "5 Cómo se conserva la generación de señales seno", 1)
    body(
        doc,
        "La compatibilidad WebGL no elimina la síntesis. La función seno sigue siendo la base matemática. Lo que cambia es "
        "que en WebGL la función vive en JavaScript y es llamada desde el proceso de audio del navegador."
    )
    code_label(doc, "Legacy en Assets scripts new Osc cs")
    code(
        doc,
        """
float SineWave(float f, int t)
{
    return Mathf.Sin(2 * Mathf.PI * f * t / FM);
}
""",
    )
    code_label(doc, "WebGL en Assets Plugins WebGL ProceduralSynth jslib")
    code(
        doc,
        """
function sine(v, f, t) {
    return Math.sin(2 * Math.PI * f * t / v.sampleRate);
}
""",
    )
    body(
        doc,
        "Esta función es equivalente a la generación seno usada por Legacy. Los parámetros son los mismos: frecuencia, tiempo "
        "de muestra y frecuencia de muestreo. Con eso se calcula la fase de la onda."
    )
    add_simple_table(
        doc,
        ["Concepto", "Legacy C#", "WebGL JavaScript"],
        [
            ("Seno", "Mathf.Sin", "Math.sin"),
            ("Pi", "Mathf.PI", "Math.PI"),
            ("Frecuencia", "f o frecuencia", "freq"),
            ("Tiempo de muestra", "TimeIndex o currentTime", "timeIndex o time"),
            ("Sample rate", "FM o AudioSettings.outputSampleRate", "v.sampleRate"),
            ("Resultado", "Valor float entre -1 y 1", "Valor number entre -1 y 1"),
        ],
        "2F5597",
    )
    body(
        doc,
        "Por eso se puede decir a los estudiantes que el proyecto no cambió de modelo sonoro: cambió de lugar de ejecución. "
        "La onda seno se sigue calculando con la misma relación entre frecuencia, tiempo y sample rate."
    )

    heading(doc, "6 Comparación del procesamiento de buffers", 1)
    body(doc, "La diferencia decisiva no está en la fórmula sino en quién solicita el siguiente bloque de muestras. En Legacy, Unity invoca OnAudioFilterRead. En WebGL, Web Audio invoca onaudioprocess sobre el nodo creado por el archivo jslib.")
    heading(doc, "6.1 Bucle Legacy dentro de OnAudioFilterRead", 2)
    code_label(doc, "Fragmento del bucle real de Assets scripts new Osc cs")
    code(doc, """
for (int i = 0; i < data.Length; i += channels)
{
    int currentTime = TimeIndex;
    float E = getADSR(currentTime);
    float X = GenerateWaveSample(FormType, currentFrequency, currentTime);
    float X1 = GenerateWaveSample(FormType, detunedFrequency, currentTime);
    float sampleValue = Level * 0.5f * (X + X1) * E * tremValue;
    data[i] = sampleValue;
    for (int ch = 1; ch < channels; ch++)
        data[i + ch] = sampleValue;
    TimeIndex++;
}
""")
    body(doc, "El índice i avanza por frame y no por canal. X y X1 son dos osciladores con detune. E es el ADSR de amplitud. Después de mezclar, el mismo valor se copia a los canales disponibles y TimeIndex avanza una sola vez.")
    heading(doc, "6.2 Bucle WebGL dentro de onaudioprocess", 2)
    code_label(doc, "Fragmento real de Assets Plugins WebGL ProceduralSynth jslib")
    code(doc, """
for (var sampleIndex = 0; sampleIndex < left.length; sampleIndex++) {
  var mixed = 0;
  for (var handleIndex = 0; handleIndex < handles.length; handleIndex++) {
    var v = voices[handles[handleIndex]];
    if (!v || !v.active || v.frequency <= 0 || v.level <= 0) continue;
    var t = v.timeIndex;
    var env = adsr(v, t);
    var currentFreq = Math.max(0, v.frequency + vib);
    var detunedFreq = currentFreq * Math.pow(2, v.detuneCents / 1200);
    mixed += v.level * 0.5 *
      (gen(v, currentFreq, t, true) + gen(v, detunedFreq, t, false)) * env * trem;
    v.timeIndex++;
  }
  mixed = Math.max(-1, Math.min(1, mixed));
  left[sampleIndex] = mixed;
  if (right !== left) right[sampleIndex] = mixed;
}
""")
    body(doc, "WebGL añade una dimensión explícita: mezcla todas las voces activas antes de escribir el frame. gen cumple el papel de GenerateWaveSample. El clamp final evita entregar al navegador valores mayores que uno o menores que menos uno.")

    heading(doc, "6.3 Selección de la forma de onda", 2)
    code_label(doc, "Despacho Legacy mediante GenerateWaveSample")
    code(doc, """
switch (type)
{
    case WaveFormType.Sine:      return SineWave(f, t);
    case WaveFormType.Square:    return SquareWave(f, t);
    case WaveFormType.Sawtooth:  return SawtoothWave(f, t);
    case WaveFormType.Triangle:  return TriangleWave(f, t);
    case WaveFormType.SA:        return SA(f, t, Narmonicos);
    case WaveFormType.FM:        return FMWave(f, t);
    case WaveFormType.Sampling1: return SamplingWave(f, t);
    case WaveFormType.WhiteNoise:return WhiteNoise(t);
    default:                     return 0f;
}
""")
    code_label(doc, "Despacho WebGL mediante gen")
    code(doc, """
if (v.useGeneratedMultiframeTimbre && v.waveform === 11)
  return generatedMultiframe(v, freq, time);
if (v.useWavetable && v.wavetable.length && v.waveform !== 5)
  return wavetable(v, time, freq);
switch (v.waveform) {
  case 0: return sine(v, freq, time);
  case 1: return sine(v, freq, time) >= 0 ? 1 : -1;
  case 4: return additive(v, freq, time);
  case 5: return fm(v, freq, time);
  case 6: return sampling(v, freq, advanceSampling);
  case 7: return (Math.random() * 2) - 1;
  default: return 0;
}
""")
    body(doc, "Los números de WebGL corresponden al orden serializado de WaveFormType. Antes del switch se priorizan el timbre generado y la wavetable válida. Esta precedencia evita que el modo generado termine en la rama default y produzca silencio.")

    heading(doc, "7 Puente entre C Sharp y JavaScript", 1)
    body(doc, "Osc conserva la responsabilidad de decidir forma de onda, ADSR y datos tímbricos. Cuando el backend externo está activo, crea un ProceduralSynthVoiceConfig. Este objeto es el contrato de datos entre la lógica de Unity y el motor que finalmente reproduce.")
    code_label(doc, "Construcción real de la configuración en Osc cs")
    code(doc, """
ProceduralSynthVoiceConfig config = new ProceduralSynthVoiceConfig
{
    sampleRate = Mathf.RoundToInt(FM),
    channels = Mathf.Max(1, streamingChannels),
    waveform = (int)FormType,
    useWavetable = ShouldReadFromWavetableInAudio(),
    level = GetEffectiveOutputLevel(),
    attackMs = A,
    decayMs = D,
    sustainMs = S,
    sustainLevel = SL,
    releaseMs = R,
    useGeneratedMultiframeTimbre = useGeneratedMultiframeTimbre,
    generatedAttackHarmonics = (float[])generatedAttackHarmonics.Clone(),
    generatedSustainHarmonics = (float[])generatedSustainHarmonics.Clone(),
    generatedReleaseHarmonics = (float[])generatedReleaseHarmonics.Clone(),
    wavetable = (float[])wavetable.Clone()
};
proceduralSynthVoice.UpdateConfig(config);
""")
    body(doc, "Las copias de los arreglos crean una instantánea estable para el backend. UpdateConfig actualiza también la voz administrada de respaldo y luego, cuando la plataforma lo permite, envía los datos al backend específico.")
    code_label(doc, "Selección real de funciones por plataforma")
    code(doc, """
#if UNITY_WEBGL && !UNITY_EDITOR
    NativeApi.PSW_SetCore(handle, config.waveform, config.level,
        config.detuneCents, config.tremLfoFrequency, config.vibLfoFrequency,
        config.vibratoDepth, config.attackMs, config.decayMs,
        config.sustainMs, config.sustainLevel, config.releaseMs);
    NativeApi.PSW_SetGeneratedTimbre(handle, enabled, attack, attackCount,
        sustain, sustainCount, release, releaseCount);
    NativeApi.PSW_SetWavetable(handle, wavetable, count);
#else
    NativeApi.PS_SetCore(handle, config.waveform, config.level,
        config.detuneCents, config.tremLfoFrequency, config.vibLfoFrequency,
        config.vibratoDepth, config.attackMs, config.decayMs,
        config.sustainMs, config.sustainLevel, config.releaseMs);
    NativeApi.PS_SetGeneratedTimbre(handle, enabled, attack, attackCount,
        sustain, sustainCount, release, releaseCount);
    NativeApi.PS_SetWavetable(handle, wavetable, count);
#endif
""")
    body(doc, "El prefijo PSW identifica las funciones WebGL expuestas por el jslib. El prefijo PS identifica el plugin externo de escritorio. El preprocesador decide qué conjunto enlazar; el resto del sintetizador usa el mismo objeto de configuración.")
    code_label(doc, "Recepción de armónicos desde memoria WebAssembly")
    code(doc, """
PSW_SetGeneratedTimbre: function(handle, enabled, attack, attackCount,
  sustain, sustainCount, release, releaseCount) {
  var v = globalThis.ProceduralSynthVoices[handle];
  if (!v) return;
  v.useGeneratedMultiframeTimbre = enabled !== 0;
  v.generatedAttackHarmonics = Array.from(
    HEAPF32.subarray(attack >> 2, (attack >> 2) + attackCount));
  v.generatedSustainHarmonics = Array.from(
    HEAPF32.subarray(sustain >> 2, (sustain >> 2) + sustainCount));
  v.generatedReleaseHarmonics = Array.from(
    HEAPF32.subarray(release >> 2, (release >> 2) + releaseCount));
}
""")
    body(doc, "Los punteros recibidos apuntan a memoria WebAssembly. El desplazamiento >> 2 convierte una dirección en bytes a un índice de float de 32 bits. Array.from crea arreglos JavaScript propios de la voz.")

    heading(doc, "8 ADSR equivalente en ambos motores", 1)
    body(doc, "Los tiempos viajan en milisegundos y se convierten a muestras mediante tiempo dividido por 1000 y multiplicado por sampleRate. Attack sube hacia uno, Decay baja hasta sustainLevel, Sustain mantiene ese nivel y Release parte del nivel real que tenía la voz al recibir NoteOff.")
    code_label(doc, "Cálculo WebGL de los límites de cada fase")
    code(doc, """
var attack = Math.max(1, Math.round((v.attackMs / 1000) * v.sampleRate));
var decay = attack + Math.round((v.decayMs / 1000) * v.sampleRate);
var sustain = decay + Math.round((v.sustainMs / 1000) * v.sampleRate);
var release = sustain + Math.round((v.releaseMs / 1000) * v.sampleRate);
if (t < attack) return curve01(t / attack, v.attackUsesLogCurve);
if (t < decay) return 1 + ((v.sustainLevel - 1) *
  curve01((t - attack) / Math.max(1, decay - attack), v.decayUsesLogCurve));
if (t < sustain) return v.sustainLevel;
if (t < release) return v.sustainLevel + ((0.0000001 - v.sustainLevel) *
  curve01((t - sustain) / Math.max(1, release - sustain), v.releaseUsesLogCurve));
return 0.0000001;
""")
    body(doc, "La versión Legacy usa los mismos campos A, D, S, SL y R al calcular getADSR. En ambos casos la envolvente se multiplica por la señal y no modifica la frecuencia. Si NoteOff ocurre antes de terminar Attack o Decay, releaseStartLevel conserva el nivel actual para evitar un salto.")
    add_simple_table(doc, ["Fase", "Entrada", "Legacy", "WebGL"], [
        ("Attack", "attackMs", "getADSR eleva E", "scheduledAdsr eleva env"),
        ("Decay", "decayMs", "E desciende a SL", "env desciende a sustainLevel"),
        ("Sustain", "sustainMs y sustainLevel", "Mantiene SL", "Mantiene v.sustainLevel"),
        ("Release", "releaseMs", "Parte de releaseStartLevel", "Parte de v.releaseStartLevel"),
    ], "404040")

    heading(doc, "9 Wavetables y lectura por fase", 1)
    body(doc, "Una wavetable almacena un ciclo periódico. El oscilador transforma la fase de la nota en un índice circular. La frecuencia controla cuántas veces por segundo se recorre la tabla.")
    code_label(doc, "Lectura administrada usada por Legacy y fallback")
    code(doc, """
int length = config.wavetable.Length;
int index = Mathf.RoundToInt((t * f / sr) * length) % length;
if (index < 0)
    index += length;
return config.wavetable[index];
""")
    code_label(doc, "Lectura equivalente en WebGL")
    code(doc, """
function wavetable(v, time, freq) {
  if (!v.wavetable.length) return 0;
  var idx = Math.round((time * freq / v.sampleRate) * v.wavetable.length)
    % v.wavetable.length;
  if (idx < 0) idx += v.wavetable.length;
  return v.wavetable[idx];
}
""")
    body(doc, "Las dos expresiones calculan el mismo índice. El módulo produce el retorno al inicio de la tabla. Esta implementación usa vecino más cercano mediante Round; mantiene equivalencia funcional, aunque una interpolación lineal sería una mejora posterior posible.")

    heading(doc, "10 Wavetables y timbres multiframe en WebGL", 1)
    body(
        doc,
        "Además de ondas básicas, WebGL también recibe wavetables y timbres generados por IA. Cuando Groq o el fallback local "
        "producen un TimbrePreset, este contiene tres conjuntos de armónicos: Attack, Sustain y Release. Estos datos se envían "
        "a WebGL para que el navegador pueda reproducir el mismo concepto tímbrico."
    )
    code(
        doc,
        """
GeneratedTimbre
  -> attackHarmonics
  -> sustainHarmonics
  -> releaseHarmonics
  -> ADSR
  -> ProceduralSynthVoiceConfig
  -> PSW_SetGeneratedTimbre
  -> generatedMultiframe en JavaScript
""",
    )
    body(
        doc,
        "En JavaScript, generatedMultiframe interpola entre los armónicos de Attack y Sustain durante el inicio de la nota. "
        "Cuando se libera la nota, interpola entre Sustain y Release. La amplitud final sigue siendo controlada por el ADSR."
    )
    body(
        doc,
        "También se conserva el antialiasing por Nyquist: para cada nota se calcula cuántos armónicos caben bajo sampleRate / 2. "
        "Los armónicos por encima de ese límite se omiten para reducir aliasing, especialmente en notas agudas."
    )

    code_label(doc, "Interpolación real en Legacy moderno")
    code(doc, """
if (releaseTriggered && currentTime >= releaseStartTimeIndex)
{
    int releaseFrames = Mathf.Max(1, Mathf.RoundToInt((R / 1000f) * FM));
    float progress = (float)(currentTime - releaseStartTimeIndex) / releaseFrames;
    HarmonicUtils.LerpHarmonics(generatedSustainHarmonics,
        generatedReleaseHarmonics, progress, generatedCurrentHarmonics);
}
else
{
    int frames = Mathf.Max(1, Mathf.RoundToInt(((A + D) / 1000f) * FM));
    HarmonicUtils.LerpHarmonics(generatedAttackHarmonics,
        generatedSustainHarmonics, (float)currentTime / frames,
        generatedCurrentHarmonics);
}
return HarmonicUtils.AdditiveSample(generatedCurrentHarmonics,
    frequency, currentTime, FM);
""")
    code_label(doc, "Síntesis multiframe real en WebGL")
    code(doc, """
var nyquistLimit = Math.floor((v.sampleRate * 0.5) / freq);
var limit = Math.min(16, nyquistLimit, from.length, to.length);
for (var h = 1; h <= limit; h++) {
  var a = Math.max(0, Math.min(1, from[h - 1] || 0));
  var b = Math.max(0, Math.min(1, to[h - 1] || 0));
  var amplitude = a + ((b - a) * progress);
  sum += Math.sin(2 * Math.PI * h * freq * time / v.sampleRate) * amplitude;
  norm += amplitude;
}
return norm > 0 ? Math.max(-1, Math.min(1, sum / norm)) : 0;
""")
    body(doc, "Cada amplitud se interpola antes de sumar su seno. norm suma las amplitudes efectivamente usadas; dividir por norm controla la salida al acumular parciales. El límite de Nyquist se calcula para cada frecuencia, por lo que una nota aguda utiliza menos armónicos que una nota grave.")
    body(doc, "Ejemplo: a 44 100 Hz, Nyquist es 22 050 Hz. Una fundamental de 440 Hz permite 50 armónicos teóricos, aunque el preset usa como máximo 16. Una fundamental de 4 000 Hz permite solo 5; el sexto ya superaría Nyquist.")

    heading(doc, "11 Ciclo de vida de una voz", 1)
    code_label(doc, "Eventos compartidos por el puente ProceduralSynthVoice")
    code(doc, """
public void NoteOn(float frequency)
{
    managedVoice.NoteOn(frequency);
    if (EnsureNativeVoice())
        NoteOnNative(nativeHandle, frequency);
}

public void NoteOff()
{
    managedVoice.NoteOff();
    if (useNativeBackend && nativeHandle >= 0)
        NoteOffNative(nativeHandle);
}
""")
    code_label(doc, "Estado reiniciado por PSW NoteOn en WebGL")
    code(doc, """
v.frequency = frequency;
v.timeIndex = 0;
v.samplingReadPosition = v.samplingStartFrame;
v.active = true;
v.releaseTriggered = false;
v.releaseStartTimeIndex = 0;
v.releaseStartLevel = 0;
""")
    body(doc, "NoteOn reinicia la fase temporal de la voz. NoteOff no destruye inmediatamente el oscilador: marca el inicio de Release para que ADSR y el frame tímbrico Release terminen suavemente. En navegador, AudioContext puede iniciar suspendido y debe reanudarse después de una interacción del usuario por la política de autoplay.")

    heading(doc, "12 Comparación integral Legacy y WebGL", 1)
    add_simple_table(
        doc,
        ["Tema", "Legacy", "WebGL"],
        [
            ("Dónde se generan muestras", "En OnAudioFilterRead dentro de Osc.cs.", "En el callback onaudioprocess del navegador dentro del jslib."),
            ("Quién pide el buffer", "Unity llama OnAudioFilterRead.", "Web Audio API solicita bloques al ScriptProcessor."),
            ("Fórmula de seno", "Mathf.Sin(2*pi*f*t/sampleRate).", "Math.sin(2*pi*f*t/sampleRate)."),
            ("ADSR", "Se calcula por muestra en C#.", "Se calcula por muestra en JavaScript con los mismos parámetros."),
            ("Wavetable", "Osc lee la tabla local o genera muestra directa.", "El jslib lee la tabla copiada desde C#."),
            ("Multiframe", "Osc interpola armónicos y genera muestra.", "generatedMultiframe interpola armónicos y genera muestra."),
            ("Estado de voz", "Campos en cada componente Osc.", "Objeto JavaScript identificado por handle."),
            ("Paso de datos", "Lectura directa de campos C#.", "ProceduralSynthVoiceConfig y funciones PSW."),
            ("Antialiasing", "HarmonicUtils limita por Nyquist.", "nyquistLimit limita antes de sumar."),
            ("Polifonía", "Varias instancias audibles de Osc.", "Tabla ProceduralSynthVoices mezclada por buffer."),
            ("Salida estéreo", "Se copia sampleValue a cada canal.", "Se escribe mixed en left y right."),
            ("Final de voz", "HasProceduralEnvelopeFinished.", "finished desactiva v.active."),
            ("Desbloqueo de audio", "No depende de política de navegador.", "Requiere interacción del usuario por reglas de autoplay."),
        ],
        "1F4E79",
    )
    body(doc, "La equivalencia buscada es musical y estructural, no identidad bit a bit. C Sharp y JavaScript pueden diferir levemente en redondeo y temporización, pero deben interpretar la misma frecuencia, el mismo preset, las mismas fases ADSR y el mismo criterio de Nyquist.")

    heading(doc, "13 Recorrido completo de una nota", 1)
    numbered(
        doc,
        [
            "El estudiante toca una nota o activa una interacción musical.",
            "Polifonia crea o reutiliza una voz Osc.",
            "Osc recibe frecuencia, forma de onda, ADSR, wavetable y timbre generado si existe.",
            "Si el backend es Legacy, Osc genera muestras en OnAudioFilterRead.",
            "Si el backend es WebGL o External DSP, Osc empaqueta la configuración en ProceduralSynthVoiceConfig.",
            "ProceduralSynthVoice envía la configuración hacia la implementación WebGL.",
            "ProceduralSynth.jslib mantiene una tabla global de voces activas.",
            "El callback de Web Audio recorre las voces, calcula seno, wavetable, ADSR o multiframe y mezcla la salida.",
            "El navegador reproduce el resultado en los canales de audio disponibles.",
        ],
    )
    body(
        doc,
        "Una frase útil para clase es: el algoritmo musical es el mismo, pero el dueño del buffer cambia. En Legacy el buffer "
        "lo entrega Unity a C#. En WebGL el buffer lo entrega el navegador a JavaScript."
    )

    code(doc, """
Interfaz -> Polifonia -> Osc
                     |-> Legacy: OnAudioFilterRead -> data[] -> AudioSource
                     |-> WebGL: ProceduralSynthVoiceConfig -> PSW_* ->
                                ProceduralSynthVoices -> onaudioprocess -> navegador
""")

    heading(doc, "14 Prácticas y pruebas para estudiantes", 1)
    add_simple_table(
        doc,
        ["Prueba", "Qué demuestra", "Resultado esperado"],
        [
            ("Sine en Legacy", "La fórmula base funciona en OnAudioFilterRead.", "Se escucha una onda seno."),
            ("Sine en WebGL", "La misma fórmula funciona en JavaScript.", "Se escucha una onda seno tras desbloquear audio."),
            ("ADSR", "La envolvente se interpreta igual.", "Attack, decay, sustain y release son audibles."),
            ("Wavetable generada", "La tabla viaja desde C# al backend.", "La forma generada suena en ambos modos."),
            ("Timbre multiframe", "Attack, Sustain y Release cambian el color.", "El sonido evoluciona durante la nota."),
            ("Nota aguda", "El límite de Nyquist reduce armónicos.", "Menos aliasing perceptible."),
            ("Cambio de backend", "La UI conserva parámetros.", "La misma configuración produce un sonido equivalente."),
        ],
        "404040",
    )

    body(doc, "Para depurar WebGL, revisar la consola del navegador y buscar los mensajes de creación de voz, NoteOn y comienzo de onaudioprocess. La compilación correcta confirma enlace y sintaxis, pero la reproducción audible debe validarse en un navegador real.")

    heading(doc, "15 Repositorio y referencias", 1)
    body(doc, "El repositorio permite comparar estos fragmentos con su contexto completo, los campos serializados y los métodos auxiliares.")
    p = doc.add_paragraph()
    p.add_run("Repositorio oficial del proyecto: ").bold = True
    add_hyperlink(p, "https://github.com/ingmiguelolivares/Audio_procedural20261", "https://github.com/ingmiguelolivares/Audio_procedural20261")
    bullets(doc, [
        "Unity Manual Audio in Web: https://docs.unity.com/en-us/engine/6000.6/manual/platform-specific/webgl/develop/audio",
        "Unity Scripting API OnAudioFilterRead: https://docs.unity.com/en-us/engine/6000.6/script-reference/unityengine/monobehaviour/onaudiofilterread",
        "Unity Manual: https://docs.unity3d.com/Manual/",
    ])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Síntesis de Audio Procedural con Inteligencia Artificial")
    r.font.size = Pt(8)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
